import pytest
import openpyxl
import io
import pandas as pd
from validator_engine import process_validation, evaluate_record, ReferenceIndex
from sample_data_generator import create_sample_files


def test_validation_engine_with_sample_data():
    mah_bytes, ref_bytes = create_sample_files()
    
    out_mah_io, qc_summary_io, stats = process_validation(mah_bytes, ref_bytes)
    
    assert stats["total_rows"] == 12
    assert stats["total_ng_found"] == 10 # 10 rows with "NG to Verify"
    
    # Check classifications
    records_by_serial = {}
    for r in stats["records"]:
        records_by_serial.setdefault(r["serial"], []).append(r)
        
    # MS001: 2 unique tickets for SAME part (Drain Pump) in ref -> Approved
    assert records_by_serial["MS001"][0]["classification"] == "Approved"
    
    # MS002: 1 unique ticket in ref (even with duplicate row) -> Non-Genuine
    assert records_by_serial["MS002"][0]["classification"] == "Non-Genuine"
    
    # MS003: Front & Rear with same ticket T3001, 0 additional tickets -> Non-Genuine
    assert len(records_by_serial["MS003"]) == 2
    assert records_by_serial["MS003"][0]["classification"] == "Non-Genuine"
    assert records_by_serial["MS003"][1]["classification"] == "Non-Genuine"
    
    # MS004: Front & Rear with T4001, plus T4002 and T4003 (2 additional suspension tickets) -> Approved
    assert len(records_by_serial["MS004"]) == 2
    assert records_by_serial["MS004"][0]["classification"] == "Approved"
    assert records_by_serial["MS004"][1]["classification"] == "Approved"
    
    # MS007: Not found in reference -> Manual Review
    assert records_by_serial["MS007"][0]["classification"] == "Manual Review"
    assert records_by_serial["MS007"][0]["status"] == "NOT_FOUND_IN_REF"
    
    # MS009: Suspension Rod with only 1 additional ticket (T9002) -> Non-Genuine
    assert records_by_serial["MS009"][0]["classification"] == "Non-Genuine"

    # MS010: Machine has 2 tickets in Ref (T1010 Valve, T1011 Pump) but evaluating Valve (only 1 ticket for Valve) -> Non-Genuine
    assert records_by_serial["MS010"][0]["classification"] == "Non-Genuine"
    
    # Verify openpyxl output preserves other rows (MS005 and MS006)
    out_wb = openpyxl.load_workbook(out_mah_io)
    ws = out_wb.active
    
    # Read rows back from processed Mahavir file
    rows = list(ws.iter_rows(values_only=True))
    header = rows[0]
    fqc_idx = header.index("FQC Analysis")
    serial_idx = header.index("Machine Serial Number")
    
    data_dict = {r[serial_idx]: r[fqc_idx] for r in rows[1:] if r[serial_idx]}
    
    assert data_dict["MS001"] == "Approved"
    assert data_dict["MS002"] == "Non-Genuine"
    assert data_dict["MS004"] == "Approved"
    assert data_dict["MS005"] == "Approved"  # preserved original
    assert data_dict["MS006"] == "DAMAGED"   # preserved original
    assert data_dict["MS007"] == "NG to Verify" # preserved original because flagged
    assert data_dict["MS010"] == "Non-Genuine"


def test_ticket_only_lookup_without_serial():
    """Verify that records with a Ticket Number present in Reference are evaluated even if Serial is missing."""
    ref_df = pd.DataFrame([
        {"Machine Serial Number": "MS999", "Ticket Number": "T9901", "Part Name": "Drain Pump"},
        {"Machine Serial Number": "MS999", "Ticket Number": "T9902", "Part Name": "Drain Pump"},
        {"Machine Serial Number": "MS888", "Ticket Number": "T8801", "Part Name": "Pulsator"}
    ])
    ref_index = ReferenceIndex(ref_df, "Machine Serial Number", "Ticket Number", "Part Name")
    
    # 1. Missing serial in Mahavir, but Ticket T9901 is in Reference (linked to MS999 with 2 tickets for Drain Pump) -> Approved
    res1 = evaluate_record(
        row_idx=1,
        serial="",
        ticket="T9901",
        part_name="Drain Pump",
        original_fqc="NG to Verify",
        ref_index=ref_index
    )
    assert res1["classification"] == "Approved"
    assert res1["status"] == "SUCCESS"
    assert res1["serial"] == "MS999"
    
    # 2. Missing serial in Mahavir, but Ticket T8801 is in Reference (linked to MS888 with 1 ticket) -> Non-Genuine
    res2 = evaluate_record(
        row_idx=2,
        serial="",
        ticket="T8801",
        part_name="Pulsator",
        original_fqc="NG to Verify",
        ref_index=ref_index
    )
    assert res2["classification"] == "Non-Genuine"
    assert res2["status"] == "SUCCESS"
    assert res2["serial"] == "MS888"
    
    # 3. Numeric ticket matching e.g. 1001.0 -> 1001
    ref_df_num = pd.DataFrame([
        {"Machine Serial Number": "100234.0", "Ticket Number": "55001.0", "Part Name": "PCB Main"}
    ])
    ref_index_num = ReferenceIndex(ref_df_num, "Machine Serial Number", "Ticket Number", "Part Name")
    res3 = evaluate_record(
        row_idx=3,
        serial="100234",
        ticket="55001",
        part_name="PCB Main",
        original_fqc="NG to Verify",
        ref_index=ref_index_num
    )
    assert res3["classification"] == "Non-Genuine"
    assert res3["status"] == "SUCCESS"


def test_normal_parts_multi_ticket_variations():
    """Test all multi-ticket combinations for normal parts (Pulsator, Pump, PCB, Valve, etc.)."""
    ref_df = pd.DataFrame([
        # Serial MS100 has 2 separate ticket rows
        {"Machine Serial Number": "MS100", "Ticket Number": "T1001", "Part Name": "Pulsator"},
        {"Machine Serial Number": "MS100", "Ticket Number": "T1002", "Part Name": "Pulsator"},
        
        # Serial MS200 has 1 ticket row, but ticket is different from Mahavir ticket
        {"Machine Serial Number": "MS200", "Ticket Number": "T2002", "Part Name": "Drain Pump"},
        
        # Serial MS300 has comma-separated tickets in single cell
        {"Machine Serial Number": "MS300", "Ticket Number": "T3001, T3002", "Part Name": "PCB Main"},
        
        # Serial MS400 has only 1 ticket across both
        {"Machine Serial Number": "MS400", "Ticket Number": "T4001", "Part Name": "Inlet Valve"}
    ])
    ref_index = ReferenceIndex(ref_df, "Machine Serial Number", "Ticket Number", "Part Name")
    
    # 1. MS100 with T1001 -> Approved (2 tickets in Ref: T1001, T1002)
    r1 = evaluate_record(1, "MS100", "T1001", "Pulsator", "NG to Verify", ref_index)
    assert r1["classification"] == "Approved"
    assert r1["status"] == "SUCCESS"
    
    # 2. MS200 with T2001 in Mahavir, T2002 in Ref -> Approved (Union has 2 distinct tickets: T2001 & T2002)
    r2 = evaluate_record(2, "MS200", "T2001", "Drain Pump", "NG to Verify", ref_index)
    assert r2["classification"] == "Approved"
    assert r2["status"] == "SUCCESS"
    assert "T2001" in r2["ref_tickets"] and "T2002" in r2["ref_tickets"]
    
    # 3. MS300 with comma separated tickets -> Approved
    r3 = evaluate_record(3, "MS300", "T3001", "PCB Main", "NG to Verify", ref_index)
    assert r3["classification"] == "Approved"
    assert r3["status"] == "SUCCESS"
    
    # 4. MS400 with same single ticket T4001 in both -> Non-Genuine
    r4 = evaluate_record(4, "MS400", "T4001", "Inlet Valve", "NG to Verify", ref_index)
    assert r4["classification"] == "Non-Genuine"
    assert r4["status"] == "SUCCESS"
    ref_index.close()


def test_persistent_reference_database_and_lookups(tmp_path):
    """Test creating persistent on-disk DuckDB database, checking metadata, and validating."""
    from validator_engine import build_persistent_reference_db, get_db_metadata
    
    test_db = str(tmp_path / "test_ref.duckdb")
    _, ref_bytes = create_sample_files()
    mah_bytes, _ = create_sample_files()
    
    # 1. Build Persistent DB
    meta = build_persistent_reference_db(
        source_input=ref_bytes,
        db_path=test_db,
        source_filename="sample_reference.xlsx"
    )
    assert meta["is_indexed"] is True
    assert meta["total_records"] > 0
    assert meta["unique_serials"] >= 6
    
    # 2. Get Metadata
    read_meta = get_db_metadata(test_db)
    assert read_meta["is_indexed"] is True
    assert read_meta["total_records"] == meta["total_records"]
    
    # 3. Process validation with ONLY Mahavir file and persistent ref_db_path (no reference_bytes)
    out_mah_io, qc_summary_io, stats = process_validation(
        mahavir_bytes=mah_bytes,
        reference_bytes=None,
        ref_db_path=test_db
    )
    assert stats["total_rows"] == 12
    assert stats["approved_count"] >= 3
    assert stats["non_genuine_count"] >= 4


def test_suspension_rod_pairing_variations_same_ticket():
    """Verify that all suspension rods sharing the same ticket number (with naming variations) are marked as pairs."""
    from validator_engine import is_suspension_rod, is_rod_or_suspension_like
    
    # 1. Test keyword coverage
    test_names = [
        "Front Suspension Rod",
        "Rear Suspension Rod",
        "Suspension Rod",
        "Front Susp Rod",
        "Rear Susp Rod",
        "Susp Rod",
        "Susp. Rod",
        "Sus Rod",
        "Sus. Rod",
        "Rod Susp",
        "Rod, Susp",
        "Rod - Susp",
        "Rod Suspension",
        "F/Suspension Rod",
        "R/Suspension Rod",
        "F-Susp Rod",
        "R-Susp Rod",
        "FR Susp Rod",
        "RR Susp Rod",
        "Suspension Spring",
        "Damper Rod",
        "Balance Rod",
        "Tub Suspension",
        "Suspension Assy",
        "Suspension Pair",
        "Suspension Set",
        "Suspension Bar",
        "SUS ASLY 6.5 SS REAR CB",
        "SUS ASLY 6.5 SS FRONT CB"
    ]
    for name in test_names:
        assert is_suspension_rod(name) is True, f"Failed to recognize suspension rod: '{name}'"
        
    # Non-suspension parts must NEVER be detected as suspension rods
    non_susp_names = [
        "INDUCTION MOTOR 45mm",
        "INDUCTION MOTOR 45mm ASSY",
        "DRAIN PUMP",
        "INLET VALVE",
        "PCB MAIN",
        "PULSATOR",
        "CLUTCH ASSEMBLY",
        "OUTER TUB",
        "MOTOR CAPACITOR",
        "PRESSURE SENSOR",
        "DOOR SWITCH"
    ]
    for name in non_susp_names:
        assert is_suspension_rod(name) is False, f"Erroneously recognized non-suspension part as suspension rod: '{name}'"
        assert is_rod_or_suspension_like(name) is False, f"Erroneously recognized non-suspension part as rod-like: '{name}'"
        
    # 2. Test Mahavir sheet with same ticket number sharing suspension parts
    mah_df = pd.DataFrame([
        {"Machine Serial Number": "MS801", "Ticket Number": "T8001", "Part Name": "Front Suspension Rod", "FQC Analysis": "NG to Verify"},
        {"Machine Serial Number": "MS801", "Ticket Number": "T8001", "Part Name": "Rear Rod", "FQC Analysis": "NG to Verify"},
        {"Machine Serial Number": "MS801", "Ticket Number": "T8001", "Part Name": "Drain Pump", "FQC Analysis": "NG to Verify"},
    ])
    
    ref_df = pd.DataFrame([
        {"Machine Serial Number": "MS801", "Ticket Number": "T8001", "Part Name": "Front Suspension Rod"},
        {"Machine Serial Number": "MS801", "Ticket Number": "T8001", "Part Name": "Rear Suspension Rod"},
        {"Machine Serial Number": "MS801", "Ticket Number": "T8002", "Part Name": "Front Suspension Rod"},
        {"Machine Serial Number": "MS801", "Ticket Number": "T8003", "Part Name": "Rear Suspension Rod"},
    ])
    
    mah_io = io.BytesIO()
    with pd.ExcelWriter(mah_io, engine="openpyxl") as writer:
        mah_df.to_excel(writer, index=False)
    mah_bytes = mah_io.getvalue()
    
    ref_io = io.BytesIO()
    with pd.ExcelWriter(ref_io, engine="openpyxl") as writer:
        ref_df.to_excel(writer, index=False)
    ref_bytes = ref_io.getvalue()
    
    _, _, stats = process_validation(mahavir_bytes=mah_bytes, reference_bytes=ref_bytes)
    
    recs = stats["records"]
    assert len(recs) == 3
    # Front Suspension Rod on T8001 -> Pair
    assert recs[0]["is_suspension"] is True
    assert recs[0]["classification"] == "Approved"
    # Rear Rod on T8001 -> Pair
    assert recs[1]["is_suspension"] is True
    assert recs[1]["classification"] == "Approved"
    # Drain Pump on T8001 -> Normal Part (only 1 ticket T8001)
    assert recs[2]["is_suspension"] is False
    assert recs[2]["classification"] == "Non-Genuine"


def test_helper_functions_exhaustive():
    """Exhaustively test utility and helper functions in validator_engine."""
    from validator_engine import (
        find_column,
        clean_val_str,
        normalize_serial,
        normalize_ticket,
        extract_tickets,
        is_ng_to_verify,
        is_suspension_rod,
        is_rod_or_suspension_like,
        SERIAL_ALIASES,
        TICKET_ALIASES,
        FQC_ALIASES,
        PART_ALIASES
    )
    
    # 1. find_column
    cols = ["SL NO", "Call ID", "FQC Remarks", "Item Desc", "Unknown"]
    assert find_column(cols, SERIAL_ALIASES) == "SL NO"
    assert find_column(cols, TICKET_ALIASES) == "Call ID"
    assert find_column(cols, FQC_ALIASES) == "FQC Remarks"
    assert find_column(cols, PART_ALIASES) == "Item Desc"
    assert find_column(cols, ["nonexistent", "dummy"]) is None
    # Substring match
    assert find_column(["my_serial_number_col"], SERIAL_ALIASES) == "my_serial_number_col"
    
    # 2. clean_val_str & normalizers
    assert clean_val_str(None) == ""
    assert clean_val_str(pd.NA) == ""
    assert clean_val_str("  Hello   World  ") == "Hello World"
    assert clean_val_str("1001.0") == "1001"
    assert clean_val_str("-500.0") == "-500"
    assert clean_val_str(12345) == "12345"
    assert normalize_serial("  MS001 ") == "MS001"
    assert normalize_serial("1001.0") == "1001"
    assert normalize_ticket(" T1001 ") == "T1001"
    assert normalize_ticket("5001.0") == "5001"
    
    # 3. extract_tickets
    assert extract_tickets("") == []
    assert extract_tickets(None) == []
    assert extract_tickets("T1001") == ["T1001"]
    assert extract_tickets("T1001, T1002/T1003; T1004\nT1005|T1006&T1007") == [
        "T1001", "T1002", "T1003", "T1004", "T1005", "T1006", "T1007"
    ]
    
    # 4. is_ng_to_verify
    assert is_ng_to_verify(None) is False
    assert is_ng_to_verify("") is False
    assert is_ng_to_verify("Approved") is False
    assert is_ng_to_verify("DAMAGED") is False
    assert is_ng_to_verify("NG to Verify") is True
    assert is_ng_to_verify("ng-to-verify") is True
    assert is_ng_to_verify("NG_TO_VERIFY") is True
    assert is_ng_to_verify("ng verification") is True
    assert is_ng_to_verify("NG") is False
    assert is_ng_to_verify("OG") is False
    assert is_ng_to_verify("DG") is False
    assert is_ng_to_verify("WRONG PART") is False
    assert is_ng_to_verify("Needs NG check") is True
    
    # 5. is_suspension_rod & is_rod_or_suspension_like
    assert is_suspension_rod(None) is False
    assert is_suspension_rod("") is False
    assert is_suspension_rod("Drain Pump") is False
    assert is_suspension_rod("Pulsator") is False
    assert is_suspension_rod("Inlet Valve") is False
    assert is_suspension_rod("PCB Main") is False
    assert is_suspension_rod("Front Suspension Rod") is True
    assert is_suspension_rod("Tub Suspension") is True
    assert is_suspension_rod("Damper Rod") is True
    assert is_suspension_rod("Balance Rod") is True
    assert is_suspension_rod("F-SUSP") is True
    assert is_suspension_rod("R-SUSP") is True
    assert is_suspension_rod("ROD/SUSP") is True
    assert is_suspension_rod("Front Rod") is True
    assert is_suspension_rod("Rear Rod") is True
    
    assert is_rod_or_suspension_like(None) is False
    assert is_rod_or_suspension_like("") is False
    assert is_rod_or_suspension_like("Drain Pump") is False
    assert is_rod_or_suspension_like("Spring") is True
    assert is_rod_or_suspension_like("Damper") is True
    assert is_rod_or_suspension_like("Rod") is True


def test_build_persistent_reference_db_all_formats(tmp_path):
    """Test build_persistent_reference_db across CSV, TSV, Parquet, and disk paths."""
    import duckdb
    from validator_engine import build_persistent_reference_db, get_db_metadata, ValidationError
    
    db_file = str(tmp_path / "master_test.duckdb")
    
    # 1. CSV bytes with comma
    csv_bytes = b"Machine Serial Number,Ticket Number,Part Name\nMS001,T1001,Front Suspension Rod\nMS001,T1002,Rear Suspension Rod\n"
    meta_csv = build_persistent_reference_db(
        source_input=csv_bytes,
        db_path=db_file,
        source_filename="ref.csv"
    )
    assert meta_csv["is_indexed"] is True
    assert meta_csv["total_records"] == 2
    
    # 2. TSV bytes with tab
    tsv_bytes = b"Machine Serial Number\tTicket Number\tPart Name\nMS002\tT2001\tDrain Pump\nMS002\tT2002\tValve\n"
    meta_tsv = build_persistent_reference_db(
        source_input=tsv_bytes,
        db_path=db_file,
        source_filename="ref.tsv"
    )
    assert meta_tsv["is_indexed"] is True
    assert meta_tsv["total_records"] == 2
    
    # 3. Parquet bytes (generated with DuckDB)
    disk_pq = str(tmp_path / "on_disk.parquet")
    con = duckdb.connect()
    con.execute("CREATE TABLE sample_pq (Machine_Serial_Number VARCHAR, Ticket_Number VARCHAR, Part_Name VARCHAR)")
    con.execute("INSERT INTO sample_pq VALUES ('MS003', 'T3001', 'PCB Main'), ('MS003', 'T3002', 'Capacitor')")
    con.execute(f"COPY sample_pq TO '{disk_pq.replace(chr(92), '/')}' (FORMAT PARQUET)")
    con.close()
    
    with open(disk_pq, "rb") as f:
        pq_bytes = f.read()
        
    meta_pq = build_persistent_reference_db(
        source_input=pq_bytes,
        db_path=db_file,
        source_filename="ref.parquet"
    )
    assert meta_pq["is_indexed"] is True
    assert meta_pq["total_records"] == 2
    
    # 4. File on disk (CSV)
    disk_csv = str(tmp_path / "on_disk.csv")
    with open(disk_csv, "w") as f:
        f.write("Machine Serial Number,Ticket Number,Part Name\nMS004,T4001,Pulsator\n")
    meta_disk_csv = build_persistent_reference_db(
        source_input=disk_csv,
        db_path=db_file,
        source_filename="on_disk.csv"
    )
    assert meta_disk_csv["is_indexed"] is True
    assert meta_disk_csv["total_records"] == 1
    
    # 5. File on disk (Parquet)
    meta_disk_pq = build_persistent_reference_db(
        source_input=disk_pq,
        db_path=db_file,
        source_filename="on_disk.parquet"
    )
    assert meta_disk_pq["is_indexed"] is True
    assert meta_disk_pq["total_records"] == 2
    
    # 6. Invalid input
    with pytest.raises(ValidationError):
        build_persistent_reference_db(
            source_input=12345,
            db_path=db_file,
            source_filename="invalid"
        )
        
    # 7. Missing required columns
    bad_csv = b"ColumnA,ColumnB\n1,2\n"
    with pytest.raises(ValidationError):
        build_persistent_reference_db(
            source_input=bad_csv,
            db_path=db_file,
            source_filename="bad.csv"
        )


def test_get_db_metadata_edge_cases(tmp_path):
    """Test get_db_metadata on non-existent, corrupt, and valid databases."""
    from validator_engine import get_db_metadata
    
    # Non-existent
    meta_none = get_db_metadata(str(tmp_path / "nonexistent.duckdb"))
    assert meta_none["is_indexed"] is False
    assert meta_none["total_records"] == 0
    
    # Corrupt file
    corrupt_file = str(tmp_path / "corrupt.duckdb")
    with open(corrupt_file, "wb") as f:
        f.write(b"NOT A DUCKDB DATABASE FILE")
    meta_corrupt = get_db_metadata(corrupt_file)
    assert meta_corrupt["is_indexed"] is False


def test_load_reference_df_and_generate_qc_workbook():
    """Test load_reference_df fallbacks and generate_qc_workbook formatting."""
    from validator_engine import load_reference_df, generate_qc_workbook, ValidationError
    
    # 1. load_reference_df with invalid bytes
    with pytest.raises(ValidationError):
        load_reference_df(b"Invalid bytes not an excel file")
        
    # 2. generate_qc_workbook
    mock_records = [
        {
            "row_idx": 2,
            "serial": "MS001",
            "ticket": "T1001",
            "part": "Drain Pump",
            "original_fqc": "NG to Verify",
            "new_fqc": "Approved",
            "classification": "Approved",
            "is_suspension": False,
            "status": "SUCCESS",
            "reason": "Rule 1 Approved",
            "ref_tickets": ["T1001", "T1002"]
        },
        {
            "row_idx": 3,
            "serial": "MS007",
            "ticket": "T7001",
            "part": "Pulsator",
            "original_fqc": "NG to Verify",
            "new_fqc": "NG to Verify",
            "classification": "Manual Review",
            "is_suspension": False,
            "status": "NOT_FOUND_IN_REF",
            "reason": "Serial 'MS007' not found in Reference File.",
            "ref_tickets": []
        }
    ]
    mock_stats = {
        "total_rows": 5,
        "total_ng_found": 2,
        "approved_count": 1,
        "non_genuine_count": 0,
        "not_found_count": 1,
        "suspension_count": 0,
        "manual_review_count": 1
    }
    qc_io = generate_qc_workbook(mock_records, mock_stats)
    assert qc_io is not None
    
    qc_wb = openpyxl.load_workbook(qc_io)
    assert "Executive Summary" in qc_wb.sheetnames
    assert "FQC Validation Log" in qc_wb.sheetnames
    assert "Manual Review Required" in qc_wb.sheetnames


def test_process_validation_validation_errors():
    """Test process_validation error paths for missing columns and invalid workbooks."""
    from validator_engine import process_validation, ValidationError
    
    ref_df = pd.DataFrame([
        {"Machine Serial Number": "MS001", "Ticket Number": "T1001", "Part Name": "Pump"}
    ])
    ref_io = io.BytesIO()
    with pd.ExcelWriter(ref_io, engine="openpyxl") as writer:
        ref_df.to_excel(writer, index=False)
    ref_bytes = ref_io.getvalue()
    
    # 1. Invalid Reference bytes
    with pytest.raises(ValidationError, match="Cannot open Reference File"):
        process_validation(mahavir_bytes=b"invalid excel", reference_bytes=b"invalid ref")

    # 2. Invalid Mahavir bytes (with valid reference bytes)
    with pytest.raises(ValidationError, match="Cannot open Mahavir File"):
        process_validation(mahavir_bytes=b"invalid excel", reference_bytes=ref_bytes)
        
    # 3. Missing FQC column in Mahavir
    bad_mah_df = pd.DataFrame([
        {"Machine Serial Number": "MS001", "Ticket Number": "T1001", "Part Name": "Pump"}
    ])
    bad_mah_io = io.BytesIO()
    with pd.ExcelWriter(bad_mah_io, engine="openpyxl") as writer:
        bad_mah_df.to_excel(writer, index=False)
    bad_mah_bytes = bad_mah_io.getvalue()
    
    with pytest.raises(ValidationError, match="missing 'FQC Analysis'"):
        process_validation(mahavir_bytes=bad_mah_bytes, reference_bytes=ref_bytes)
        
    # 4. Missing Serial column in Mahavir
    no_serial_df = pd.DataFrame([
        {"FQC Analysis": "NG to Verify", "Ticket Number": "T1001", "Part Name": "Pump"}
    ])
    no_serial_io = io.BytesIO()
    with pd.ExcelWriter(no_serial_io, engine="openpyxl") as writer:
        no_serial_df.to_excel(writer, index=False)
        
    with pytest.raises(ValidationError, match="missing 'Machine Serial Number'"):
        process_validation(mahavir_bytes=no_serial_io.getvalue(), reference_bytes=ref_bytes)
        
    # 5. No reference provided and non-existent DB path
    with pytest.raises(ValidationError, match="No Reference File uploaded"):
        process_validation(mahavir_bytes=bad_mah_bytes, reference_bytes=None, ref_db_path="nonexistent.duckdb")


def test_evaluate_record_edge_cases():
    """Test evaluate_record edge cases: missing identifiers, empty reference lookups."""
    from validator_engine import evaluate_record, ReferenceIndex
    
    ref_df = pd.DataFrame([
        {"Machine Serial Number": "MS101", "Ticket Number": "", "Part Name": "Pump"}
    ])
    ref_idx = ReferenceIndex(ref_df, "Machine Serial Number", "Ticket Number", "Part Name")
    
    # 1. Both serial and ticket empty
    rec1 = evaluate_record(1, "", "", "Pump", "NG to Verify", ref_idx)
    assert rec1["classification"] == "Manual Review"
    assert rec1["status"] == "FLAGGED"
    
    # 2. Serial exists in ref but has 0 tickets
    rec2 = evaluate_record(2, "MS101", "", "Pump", "NG to Verify", ref_idx)
    assert rec2["classification"] == "Non-Genuine"
    
    # 3. Double close
    ref_idx.close()
    ref_idx.close()


def test_build_persistent_reference_db_from_dataframe(tmp_path):
    """Test building persistent DB directly from a pandas DataFrame."""
    from validator_engine import build_persistent_reference_db, get_db_metadata
    
    df = pd.DataFrame([
        {"Machine Serial Number": "MS901", "Ticket Number": "T9001", "Part Name": "Front Susp Rod"},
        {"Machine Serial Number": "MS901", "Ticket Number": "T9002", "Part Name": "Rear Susp Rod"},
        {"Machine Serial Number": "MS902", "Ticket Number": "", "Part Name": "Valve"}
    ])
    db_file = str(tmp_path / "df_test.duckdb")
    meta = build_persistent_reference_db(df, db_file, "df_source")
    assert meta["is_indexed"] is True
    assert meta["total_records"] == 3


def test_process_validation_empty_db_and_consecutive_empty_rows(tmp_path):
    """Test validation with empty DB and files with empty trailing rows."""
    from validator_engine import process_validation, ValidationError, get_db_metadata
    import duckdb
    
    # 1. Persistent DB with 0 records
    empty_db = str(tmp_path / "empty.duckdb")
    con = duckdb.connect(empty_db)
    con.execute("CREATE TABLE db_metadata (indexed_at VARCHAR, source_filename VARCHAR, total_records INTEGER, unique_serials INTEGER, unique_tickets INTEGER)")
    con.execute("INSERT INTO db_metadata VALUES ('now', 'empty', 0, 0, 0)")
    con.close()
    
    mah_bytes, _ = create_sample_files()
    with pytest.raises(ValidationError, match="contains no records"):
        process_validation(mahavir_bytes=mah_bytes, ref_db_path=empty_db)

    # 2. Mahavir file with 120 trailing blank rows
    mah_df = pd.DataFrame([
        {"Machine Serial Number": "MS001", "Ticket Number": "T1001", "Part Name": "Pump", "FQC Analysis": "NG to Verify"}
    ] + [{"Machine Serial Number": None, "Ticket Number": None, "Part Name": None, "FQC Analysis": None}] * 120)
    
    mah_io = io.BytesIO()
    with pd.ExcelWriter(mah_io, engine="openpyxl") as writer:
        mah_df.to_excel(writer, index=False)
        
    ref_df = pd.DataFrame([
        {"Machine Serial Number": "MS001", "Ticket Number": "T1001", "Part Name": "Pump"},
        {"Machine Serial Number": "MS001", "Ticket Number": "T1002", "Part Name": "Pump"}
    ])
    ref_io = io.BytesIO()
    with pd.ExcelWriter(ref_io, engine="openpyxl") as writer:
        ref_df.to_excel(writer, index=False)
        
    _, _, stats = process_validation(mah_io.getvalue(), ref_io.getvalue())
    assert stats["total_rows"] == 1
    assert stats["approved_count"] == 1


def test_same_part_vs_different_part_multi_ticket():
    """Explicitly verify that multiple tickets must be for the SAME part to approve."""
    ref_df = pd.DataFrame([
        # Serial MS500 has 2 tickets for Drain Pump
        {"Machine Serial Number": "MS500", "Ticket Number": "T5001", "Part Name": "Drain Pump"},
        {"Machine Serial Number": "MS500", "Ticket Number": "T5002", "Part Name": "Drain Pump"},
        
        # Serial MS600 has 2 tickets for DIFFERENT parts
        {"Machine Serial Number": "MS600", "Ticket Number": "T6001", "Part Name": "Drain Pump"},
        {"Machine Serial Number": "MS600", "Ticket Number": "T6002", "Part Name": "Inlet Valve"},
        
        # Serial MS700 has token reordered part names: 'PCB Main' and 'Main PCB'
        {"Machine Serial Number": "MS700", "Ticket Number": "T7001", "Part Name": "PCB Main"},
        {"Machine Serial Number": "MS700", "Ticket Number": "T7002", "Part Name": "Main PCB"},
    ])
    ref_idx = ReferenceIndex(ref_df, "Machine Serial Number", "Ticket Number", "Part Name")
    
    # 1. MS500 evaluated for Drain Pump -> Approved (2 tickets for Drain Pump)
    r1 = evaluate_record(1, "MS500", "T5001", "Drain Pump", "NG to Verify", ref_idx)
    assert r1["classification"] == "Approved"
    assert "2 distinct ticket(s) for part 'Drain Pump'" in r1["reason"]
    
    # 2. MS500 evaluated for Inlet Valve with T5001 -> Non-Genuine (only 1 ticket for Inlet Valve)
    r2 = evaluate_record(2, "MS500", "T5001", "Inlet Valve", "NG to Verify", ref_idx)
    assert r2["classification"] == "Non-Genuine"
    assert "Only 1 unique ticket found for part 'Inlet Valve'" in r2["reason"]
    
    # 3. MS500 evaluated for Inlet Valve without ticket -> Non-Genuine (0 tickets for Inlet Valve in Ref)
    r2_no_tkt = evaluate_record(2, "MS500", "", "Inlet Valve", "NG to Verify", ref_idx)
    assert r2_no_tkt["classification"] == "Non-Genuine"
    assert "0 tickets for part 'Inlet Valve'" in r2_no_tkt["reason"]
    
    # 4. MS600 evaluated for Drain Pump -> Non-Genuine (only 1 ticket for Drain Pump)
    r3 = evaluate_record(3, "MS600", "T6001", "Drain Pump", "NG to Verify", ref_idx)
    assert r3["classification"] == "Non-Genuine"
    assert "Only 1 unique ticket found for part 'Drain Pump'" in r3["reason"]
    
    # 5. MS700 evaluated for PCB Main -> Approved (token reordering matched)
    r4 = evaluate_record(4, "MS700", "T7001", "PCB Main", "NG to Verify", ref_idx)
    assert r4["classification"] == "Approved"
    assert "2 distinct ticket(s) for part 'PCB Main'" in r4["reason"]
    
    ref_idx.close()


def test_load_reference_df_multisheet_offset_header():
    """Test loading multi-sheet Excel where headers are offset."""
    from validator_engine import load_reference_df
    
    wb = openpyxl.Workbook()
    # Sheet 1: Empty rows before headers
    ws1 = wb.active
    ws1.title = "Sheet1"
    ws1.append(["Title", "Report"])
    ws1.append([])
    ws1.append(["Machine Serial Number", "Ticket Number", "Part Name"])
    ws1.append(["MS501", "T5001", "Pulsator"])
    
    # Sheet 2: Standard
    ws2 = wb.create_sheet(title="Sheet2")
    ws2.append(["Machine Serial Number", "Ticket Number", "Part Name"])
    ws2.append(["MS502", "T5002", "Inlet Valve"])
    
    bio = io.BytesIO()
    wb.save(bio)
    ref_df, s_col, t_col, p_col = load_reference_df(bio.getvalue())
    assert len(ref_df) == 2
    assert s_col == "Machine Serial Number"


def test_offset_header_and_reindex(tmp_path):
    from validator_engine import build_persistent_reference_db
    wb = openpyxl.Workbook()
    ws = wb.active
    # Empty header offset rows and empty leading columns
    ws.append([])
    ws.append(["Title Info", "", "", ""])
    ws.append([])
    # Row 4: Column A is None, Column B is None, Column C is Serial, Column D is Ticket, Column E is Part
    ws.append([None, None, "Machine Serial Number", "Ticket Number", "Part Name"])
    # Row 5 & 6 data with leading None
    ws.append([None, None, "MS_OFFSET_101", "T_OFF_001", "Drain Pump"])
    ws.append([None, None, "MS_OFFSET_101", "T_OFF_002", "Drain Pump"])
    
    bio = io.BytesIO()
    wb.save(bio)
    excel_bytes = bio.getvalue()
    
    db_file = str(tmp_path / "offset_test.duckdb")
    
    # 1. First indexing run
    meta1 = build_persistent_reference_db(excel_bytes, db_file, "offset_reference.xlsx")
    assert meta1["is_indexed"] is True
    assert meta1["total_records"] == 2
    assert meta1["unique_serials"] == 1
    
    # 2. Re-indexing run on same file/path (must not throw 'Table already exists')
    meta2 = build_persistent_reference_db(excel_bytes, db_file, "offset_reference.xlsx")
    assert meta2["is_indexed"] is True
    assert meta2["total_records"] == 2
    
    # 3. Evaluate record
    ref_idx = ReferenceIndex(db_path=db_file)
    res = evaluate_record(1, "MS_OFFSET_101", "T_OFF_001", "Drain Pump", "NG to Verify", ref_idx)
    assert res["classification"] == "Approved"
    ref_idx.close()


def test_concatenated_suspension_rod_codes(tmp_path):
    from validator_engine import build_persistent_reference_db
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.append(["Machine Serial Number", "Ticket Number", "Part Description"])
    ws.append(["021008180301118403", "1015644848", "SUSASLY6.5SSREARCBWG"])
    ws.append(["021008180301118403", "1015644848", "SUSASLY6.5SSFRONTCBWG"])
    ws.append(["021008180301118403", "1017652172", "SUSASLY6.5SSFRONTCBWG"])
    ws.append(["021008180301118403", "1017652172", "SUSASLY6.5SSREARCBWG"])

    bio = io.BytesIO()
    wb.save(bio)
    excel_bytes = bio.getvalue()
    db_file = str(tmp_path / "concat_susp.duckdb")

    meta = build_persistent_reference_db(excel_bytes, db_file, "concat_ref.xlsx")
    assert meta["is_indexed"] is True
    
    ref_idx = ReferenceIndex(db_path=db_file)
    # Claim from Mahavir with spaced name and new ticket 1020260179
    res = evaluate_record(102, "021008180301118403", "1020260179", "SUS ASLY 6.5 SS REAR CB ...", "ng to verify", ref_idx)
    assert res["is_suspension"] is True
    assert res["classification"] == "Approved"
    assert len(res["additional_tickets"]) == 2
    assert set(res["additional_tickets"]) == {"1015644848", "1017652172"}
    ref_idx.close()


def test_multisheet_mahavir_validation(tmp_path):
    from validator_engine import process_validation, build_persistent_reference_db
    
    # 1. Create Reference workbook with entries for all appliances
    wb_ref = openpyxl.Workbook()
    ws_ref = wb_ref.active
    ws_ref.append(["Machine Serial Number", "Ticket Number", "Part Description"])
    # TL: 2 historical suspension tickets
    ws_ref.append(["MS_TL_01", "T_TL_H1", "SUSASLY6.5SSREARCBWG"])
    ws_ref.append(["MS_TL_01", "T_TL_H2", "SUSASLY6.5SSFRONTCBWG"])
    # FL: 2 historical pump tickets
    ws_ref.append(["MS_FL_01", "T_FL_H1", "Drain Pump"])
    ws_ref.append(["MS_FL_01", "T_FL_H2", "Drain Pump"])
    # MW: 1 ticket only (Non-genuine)
    ws_ref.append(["MS_MW_01", "T_MW_01", "Magnetron"])
    # DW: unindexed serial (Manual review)
    bio_ref = io.BytesIO()
    wb_ref.save(bio_ref)
    ref_bytes = bio_ref.getvalue()
    
    # 2. Create Mahavir workbook with sheets TL, FL, MW, DW
    wb_mah = openpyxl.Workbook()
    
    # Sheet TL (Top Load)
    ws_tl = wb_mah.active
    ws_tl.title = "TL"
    ws_tl.append(["Machine Serial Number", "Ticket Number", "Part Description", "FQC Analysis"])
    ws_tl.append(["MS_TL_01", "T_TL_CLAIM", "SUS ASLY 6.5 SS REAR", "ng to verify"])
    
    # Sheet FL (Front Load / UF)
    ws_fl = wb_mah.create_sheet(title="FL")
    ws_fl.append(["Machine Serial Number", "Ticket Number", "Part Description", "FQC Analysis"])
    ws_fl.append(["MS_FL_01", "T_FL_CLAIM", "Drain Pump", "NG to Verify"])
    
    # Sheet MW (Microwave / MV)
    ws_mw = wb_mah.create_sheet(title="MW")
    ws_mw.append(["Machine Serial Number", "Ticket Number", "Part Description", "FQC Analysis"])
    ws_mw.append(["MS_MW_01", "T_MW_01", "Magnetron", "NG TO VERIFY"])
    
    # Sheet DW (Dishwasher / DV)
    ws_dw = wb_mah.create_sheet(title="DW")
    ws_dw.append(["Machine Serial Number", "Ticket Number", "Part Description", "FQC Analysis"])
    ws_dw.append(["MS_DW_99", "T_DW_99", "Inlet Valve", "NG to Verify"])
    
    bio_mah = io.BytesIO()
    wb_mah.save(bio_mah)
    mah_bytes = bio_mah.getvalue()
    
    # 3. Process Validation
    out_mah_io, qc_sum_io, stats = process_validation(mah_bytes, reference_bytes=ref_bytes)
    
    assert stats["processed_sheets"] == ["TL", "FL", "MW", "DW"]
    assert stats["total_rows"] == 4
    assert stats["total_ng_found"] == 4
    assert stats["approved_count"] == 2      # TL (suspension >= 2) and FL (drain pump >= 2)
    assert stats["non_genuine_count"] == 1   # MW (only 1 ticket for magnetron)
    assert stats["manual_review_count"] == 1 # DW (unindexed serial)
    
    # Verify sheet names are attached to evaluated records
    sheet_names_found = [r["sheet_name"] for r in stats["records"]]
    assert sheet_names_found == ["TL", "FL", "MW", "DW"]
    
    # Verify all 4 sheets in output Mahavir workbook were updated in-place
    out_wb = openpyxl.load_workbook(out_mah_io)
    assert out_wb.sheetnames == ["TL", "FL", "MW", "DW"]
    assert out_wb["TL"].cell(row=2, column=4).value == "Approved"
    assert out_wb["FL"].cell(row=2, column=4).value == "Approved"
    assert out_wb["MW"].cell(row=2, column=4).value == "Non-Genuine"
    assert out_wb["DW"].cell(row=2, column=4).value == "NG to Verify" # Unchanged for manual review


if __name__ == "__main__":
    pytest.main(["-v", "backend/test_validator.py"])








